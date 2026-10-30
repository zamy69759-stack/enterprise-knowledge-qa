#!/usr/bin/env python3
"""Resolve explicit knowledge-image associations and stage real images for display."""
import argparse, hashlib, io, json, re, shutil, sys
from pathlib import Path, PureWindowsPath
from urllib.parse import unquote
from PIL import Image
import kb

BAD = ('已废弃', '作废', '不参与关联', '非业务元数据', 'macos appledouble')

def discarded(record):
    return any(t in kb.norm(' '.join(str(v or '') for v in record.values())) for t in BAD)

# Column aliases describe relationships, not a specific product or workbook.
ALIASES = {
 'entry_id': ('知识编号','条目编号','记录编号','案例编号','步骤编号','知识ID','条目ID','id','qa_id','entry_id','record_id'),
 'image_id': ('图片编号','图片ID','图像编号','image_id','imageid','asset_id'),
 'linked_ids': ('关联图片编号','关联图片ID','图片编号','图片ID','image_ids','image_id','linked_image_ids','images'),
 'path': ('图片相对路径','图片路径','图像路径','文件路径','图片/附件引用','图片引用','附件引用','image_path','image_paths','path','file_path','src','href','url'),
 'source': ('原始来源','来源文档','来源文件','源文件','文档路径','source','source_file','source_path','document','document_path'),
 'review': ('审核状态','状态','review_status','status'),
 'description': ('图片描述','图片说明','图注','标题','图片名称','description','caption','alt','title'),
}
def field(rec, kind, default=None):
    keys={kb.norm(str(k).strip()):v for k,v in rec.items()}
    for key in ALIASES[kind]:
        value=keys.get(kb.norm(key))
        if value is not None and value!='': return value
    return default

def records(ws):
    # Locate a recognizable header even when a sheet begins with title/metadata rows.
    rows=iter(ws.iter_rows(values_only=True)); header=(); start=1
    recognized={kb.norm(a) for v in ALIASES.values() for a in v}
    for n,row in enumerate(rows,1):
        if any(kb.norm(c) in recognized for c in row if c is not None):
            header=row;start=n;break
        if n>=20: return
    for n,row in enumerate(rows,start+1):
        yield n,{str(k).strip():v for k,v in zip(header,row) if k is not None}

def ids(value):
    if isinstance(value,(list,tuple)): return [str(v).strip() for v in value if str(v).strip()]
    # IDs may be GUIDs, numeric strings, non-English codes, or custom identifiers.
    return [v.strip() for v in re.split(r'[,，;；\n]+',str(value or '')) if v.strip()]

def normalize_record(rec):
    return {**rec,'审核状态':field(rec,'review','未注明'),'图片描述':field(rec,'description',''),
            '图片编号':field(rec,'image_id'),'图片相对路径':field(rec,'path'),'原始来源':field(rec,'source','')}

def paths(value):
    s=str(value or '').strip()
    # A path can contain spaces. Split only explicit multi-value separators.
    return [v.strip().strip('"') for v in re.split(r'[\n;；]+',s) if v.strip()]

def linked_file(roots,base,raw):
    raw=unquote(raw.strip().strip('<>')).replace('\\','/')
    if re.match(r'^[a-zA-Z][a-zA-Z0-9+.-]*:',raw) or raw.startswith('//') or PureWindowsPath(raw).is_absolute() or Path(raw).is_absolute():
        raise ValueError('拒绝外链、绝对路径或网络路径引用')
    p=base/Path(raw)
    # Permit ../images only if final location remains inside configured roots.
    matches=[r for r in roots.values() if p.resolve().is_relative_to(r.resolve())]
    if not matches: raise ValueError('图片引用超出已配置知识库')
    root=matches[0]
    for segment in (p,*p.parents):
        if segment==root: break
        if segment.is_symlink(): raise ValueError('拒绝符号链接图片')
    p=p.resolve()
    if not p.is_file(): raise ValueError('关联图片文件不存在或不可访问')
    if p.stat().st_size>50*1024*1024: raise ValueError('图片大于50MiB')
    return p

def collect(roots,source,args):
    candidates,issues=[],[]
    def add(raw,origin,record=None):
        record=record or {}
        if discarded(record): return
        candidates.append({'reference':raw,'origin':origin,'review_status':record.get('审核状态','未注明'),
                           'association_status':record.get('图片关联审核状态',record.get('审核状态','未注明')),
                           'association_type':record.get('_association_type','row-reference'),'image_id':record.get('图片编号'),'description':record.get('图片描述',record.get('图片名称',''))})
    if args.reference:
        for raw in args.reference:
            add(raw,args.provenance or '由Codex从当前来源文件读取的显式图片引用')
    if source.suffix.lower()=='.xlsx':
        import openpyxl
        wb=openpyxl.load_workbook(source,read_only=True,data_only=True)
        try:
            if args.sheet and args.sheet not in wb.sheetnames: raise ValueError('没有指定的知识工作表')
            all_sheets={ws.title:list(records(ws)) for ws in wb}
            image_sheets=set()
            index={}
            for title,rows in all_sheets.items():
                for n,rec in rows:
                    # Image indexes carry an image identifier/path, without a knowledge entry ID.
                    if field(rec,'image_id') and field(rec,'path') and not field(rec,'entry_id'):
                        image_sheets.add(title)
                        image=normalize_record(rec);key=str(field(rec,'image_id'))
                        index.setdefault(key,[]).append((title,n,image))
            matched=0;source_groups=set()
            for title,rows in all_sheets.items():
                if args.sheet and title!=args.sheet: continue
                if title in image_sheets: continue
                for row,rec in rows:
                    if row not in args.row and str(field(rec,'entry_id')) not in args.qa_id: continue
                    matched+=1
                    if discarded(rec): continue
                    source_group=str(field(rec,'source','')).split('｜')[0].strip()
                    if source_group: source_groups.add(source_group)
                    origin=f'工作表《{title}》第{row}行，{field(rec,"entry_id", "")}'
                    for raw in paths(field(rec,'path')): add(raw,origin,normalize_record(rec))
                    for key in ids(field(rec,'linked_ids')):
                        if key not in index: issues.append({'origin':origin,'image_id':key,'error':'当前工作簿未找到该图片编号；请检查外部图片索引或用--reference传入已核实的图片引用'})
                        for ititle,n,image in index.get(key,[]):
                            for raw in paths(field(image,'path')):
                                combined={**image,'图片关联审核状态':rec.get('图片关联审核状态') or field(image,'review')}
                                add(raw,origin+f'；《{ititle}》第{n}行',combined)
            for key,entries in index.items():
                for title,n,image in entries:
                    group=str(field(image,'source','')).split('｜')[0].strip()
                    if (group and group in source_groups) or key in args.image_id:
                        kind='selected-index-id' if key in args.image_id else 'same-source-candidate'
                        for raw in paths(field(image,'path')):
                            add(raw,f'《{title}》第{n}行；来源：{group}',{**image,'_association_type':kind})
            if not matched and not args.image_id and not args.reference: issues.append({'error':'未找到所选条目，需检查表头、工作表或条目ID；可用--sheet和--row定位，或读取原资料后用--reference传入显式引用'})
        finally: wb.close()
    elif source.suffix.lower()=='.md':
        lines=kb.decode(source.read_bytes()).splitlines(); references={}
        for line in lines:
            m=re.match(r'^\s*\[([^]]+)\]:\s*<?([^>\s]+)>?',line)
            if m: references[m[1].casefold()]=m[2]
        for n,line in enumerate(lines,1):
            if not args.start<=n<=args.end: continue
            for m in re.finditer(r'!\[([^]]*)\]\(\s*(<[^>]+>|[^)]+)\)',line):
                raw=m[2]
                if raw.startswith('<'): raw=raw[1:raw.find('>')]
                else: raw=re.sub(r'\s+[\"\'][^\"\']*[\"\']\s*$','',raw)
                add(raw,f'MD第{n}行',{'图片描述':m[1]})
            for m in re.finditer(r'!\[([^]]*)\]\[([^]]*)\]',line):
                key=(m[2] or m[1]).casefold()
                if key in references: add(references[key],f'MD第{n}行',{'图片描述':m[1]})
    elif not args.reference: raise ValueError('该格式需先读取原文档或索引，再通过--reference和--provenance传入已核实的关联路径；不得按文件名猜测')
    return candidates,issues

def stage(roots,source,args):
    out=Path(args.output_dir).expanduser().resolve()
    if any(out.is_relative_to(r.resolve()) for r in roots.values()): raise ValueError('展示副本必须放在本机会话目录，不能写入知识库')
    candidates,issues=collect(roots,source,args); valid=[]; seen=set()
    for item in candidates:
        try:
            p=linked_file(roots,source.parent,item['reference']); data=p.read_bytes()
            with Image.open(io.BytesIO(data)) as im:
                width,height=im.size; fmt=im.format
                if width*height>60_000_000: raise ValueError('图片像素过大，需专用查看器')
                im.verify()
            sha=hashlib.sha256(data).hexdigest()
            if sha in seen: continue
            seen.add(sha)
            valid.append({**item,'source_path':str(p),'sha256':sha,'width':width,'height':height,'format':fmt})
        except Exception as e: issues.append({**item,'error':str(e)})
    out.mkdir(parents=True,exist_ok=True)
    selected=valid[:args.limit]
    for item in selected:
        ext={'PNG':'.png','JPEG':'.jpg','WEBP':'.webp','GIF':'.gif'}.get(item['format'])
        target=out/(item['sha256'][:24]+(ext or '.png'))
        if not target.exists():
            if ext: shutil.copyfile(item['source_path'],target)
            else:
                with Image.open(item['source_path']) as im: im.save(target,format='PNG')
        else:
            with Image.open(target) as im: im.verify()
        item['display_path']=str(target)
        item['markdown_example']=f'![关联图片，须先查看核对内容]({target.as_posix()})'
    return {'source':str(source),'images':selected,'total_valid_images':len(valid),'truncated':len(valid)>len(selected),'issues':issues,
            'notice':'输出仅确认显式关联及文件可读。必须查看图片核对内容、审核状态与回答相关性，再实际在最终回复内展示。脚本成功不代表客户端已显示。'}

def main():
    a=argparse.ArgumentParser();a.add_argument('--config');a.add_argument('--category',required=True,choices=kb.CATS);a.add_argument('--file',required=True)
    a.add_argument('--sheet',help='可选；自动识别知识及图片工作表，不固定工作表名称');a.add_argument('--qa-id',action='append',default=[]);a.add_argument('--image-id',action='append',default=[]);a.add_argument('--row',type=int,action='append',default=[])
    a.add_argument('--reference',action='append',default=[],help='从源文档或索引核实的显式相对图片路径');a.add_argument('--provenance',help='显式引用在源文档/索引中的定位');
    a.add_argument('--start',type=int,default=1);a.add_argument('--end',type=int,default=200);a.add_argument('--limit',type=int,default=6);a.add_argument('--output-dir',required=True)
    args=a.parse_args()
    if not 1<=args.limit<=20: raise ValueError('limit须为1到20')
    roots=kb.load_config(kb.config_path(args.config));source=kb.resolve(roots,args.category,args.file)
    if source.suffix.lower()=='.xlsx' and not (args.qa_id or args.row or args.image_id or args.reference): raise ValueError('请指定相关知识编号--qa-id或Excel实际行号--row，禁止无选择导出全表图片')
    if not 1<=args.start<=args.end: raise ValueError('MD行号范围无效')
    kb.emit(stage(roots,source,args))
if __name__=='__main__':
    if hasattr(sys.stdout,'reconfigure'):sys.stdout.reconfigure(encoding='utf-8')
    try:main()
    except Exception as e:kb.emit({'status':'error','error':str(e),'notice':'配图未完成，不得声称已展示图片。'});sys.exit(2)
