#!/usr/bin/env python3
"""Read-only evidence retrieval; Codex performs semantic expansion and answering."""
import argparse, csv, io, json, os, sys, unicodedata
from pathlib import Path
CATS = ('企业基础', '产品知识', '企业Q&A', '业务SOP')
EXTS = {'.md', '.txt', '.csv', '.xlsx', '.docx', '.pptx', '.pdf'}
def norm(s): return unicodedata.normalize('NFKC',str(s)).casefold()
def emit(x): print(json.dumps(x,ensure_ascii=False,indent=2,default=str))
def config_path(s): return Path(s or os.getenv('ENTERPRISE_QA_CONFIG') or Path.home()/'.enterprise-knowledge-qa'/'config.json').expanduser()
def load_config(p):
    d=json.loads(p.read_text(encoding='utf-8-sig'))['folders']
    if set(d)!=set(CATS): raise ValueError('folders必须且只能包含四个指定分类')
    if any(not isinstance(v,str) or not Path(v).is_absolute() for v in d.values()): raise ValueError('分类路径必须是当前环境可访问的绝对路径')
    return {k:Path(v) for k,v in d.items()}
def status(roots):
    out={}
    for k,p in roots.items():
        try:
            with os.scandir(p): pass
            out[k]={'path':str(p),'status':'accessible'}
        except OSError as e: out[k]={'path':str(p),'status':'unavailable','error':str(e)}
    return out

def inventory(roots):
    files, warnings=[],[]; states=status(roots)
    for cat,root in roots.items():
        if states[cat]['status']!='accessible':
            warnings.append({'category':cat,'error':states[cat]['error']}); continue
        def fail(e): warnings.append({'category':cat,'error':str(e)})
        for current,dirs,names in os.walk(root,onerror=fail,followlinks=False):
            dirs[:]=sorted(d for d in dirs if not d.startswith('.') and not Path(current,d).is_symlink())
            for name in sorted(names):
                p=Path(current,name)
                if name.startswith(('.','~$')) or p.is_symlink(): continue
                try:
                    p.resolve().relative_to(root.resolve()); st=p.stat()
                    entry={'category':cat,'file':p.relative_to(root).as_posix(),'bytes':st.st_size,'mtime_ns':st.st_mtime_ns,'supported':p.suffix.lower() in EXTS}
                    files.append(entry)
                    if not entry['supported']: warnings.append({**entry,'warning':'不支持该格式；正文尚未检索，必要时须专用工具读取'})
                except (OSError,ValueError) as e: fail(e)
    return files,warnings,states

def resolve(roots,cat,relative):
    root=roots[cat]; p=root/relative
    p.resolve().relative_to(root.resolve())
    if Path(relative).is_absolute() or '..' in Path(relative).parts: raise ValueError('请使用分类内相对路径')
    probe=p
    while probe!=root:
        if probe.is_symlink(): raise ValueError('不读取符号链接')
        if probe.parent==probe: raise ValueError('路径越界')
        probe=probe.parent
    if not p.is_file(): raise ValueError('文件不存在或不可读')
    if p.stat().st_size>50*1024*1024: raise ValueError('文件超过50MiB，需分段读取或专用工具')
    return p

def decode(b):
    if b.startswith((b'\xff\xfe',b'\xfe\xff')): return b.decode('utf-16')
    for enc in ('utf-8-sig','gb18030'):
        try: return b.decode(enc)
        except UnicodeError: pass
    raise ValueError('无法可靠识别文本编码')

def extract(p):
    units,notes=[],[]
    def add(loc,text): units.append({'locator':loc,'text':str(text)})
    ext=p.suffix.lower()
    if ext in ('.md','.txt'):
        for i,line in enumerate(decode(p.read_bytes()).splitlines(),1): add(f'行{i}',line)
    elif ext=='.csv':
        for i,row in enumerate(csv.reader(io.StringIO(decode(p.read_bytes()))),1): add(f'CSV记录{i}',json.dumps(row,ensure_ascii=False))
    elif ext=='.xlsx':
        import openpyxl
        from openpyxl.utils import get_column_letter
        wb=openpyxl.load_workbook(p,read_only=True,data_only=False)
        values=openpyxl.load_workbook(p,read_only=True,data_only=True)
        try:
            for ws in wb:
                if ws.sheet_state!='visible': notes.append(f'工作表《{ws.title}》为隐藏状态，仍属于该可访问文件')
                for row,cached in zip(ws.iter_rows(),values[ws.title].iter_rows()):
                    parts=[]
                    for col,(c,v) in enumerate(zip(row,cached),1):
                        if c.value is None: continue
                        text=str(c.value)
                        if c.data_type=='f': text+=f' [公式缓存={v.value}；未重算]' if v.value is not None else ' [无公式缓存结果，不能作为确认数值]'
                        parts.append(f'{get_column_letter(col)}={text}')
                    if parts: add(f'工作表《{ws.title}》第{row[0].row}行',' | '.join(parts))
            notes.append('已读取Excel单元格正文，嵌入图片、图中文字、超链接目标未自动识别，必要时另行核实。')
        finally: wb.close(); values.close()
    elif ext=='.docx':
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        doc=Document(p)
        for i,el in enumerate(doc.element.body.iterchildren(),1):
            if el.tag==qn('w:p'): add(f'正文块{i}',Paragraph(el,doc).text)
            elif el.tag==qn('w:tbl'):
                for j,row in enumerate(Table(el,doc).rows,1): add(f'正文块{i}表格行{j}',' | '.join(c.text for c in row.cells))
        notes.append('已读取Word正文和表格；图片、文本框、页眉页脚未保证覆盖。')
    elif ext=='.pptx':
        from pptx import Presentation
        def flatten(shapes):
            for sh in shapes:
                if sh.shape_type==6: yield from flatten(sh.shapes)
                else: yield sh
        for i,slide in enumerate(Presentation(p).slides,1):
            for j,sh in enumerate(flatten(slide.shapes),1):
                if sh.has_text_frame: add(f'第{i}页文本{j}',sh.text)
                if sh.has_table:
                    for k,row in enumerate(sh.table.rows,1): add(f'第{i}页表格{j}行{k}',' | '.join(c.text for c in row.cells))
        notes.append('已读取PPT可编辑文本和表格；图片和图表内容未自动识别。')
    elif ext=='.pdf':
        import pymupdf
        with pymupdf.open(p) as doc:
            for i,page in enumerate(doc,1):
                content=page.get_text(sort=True); add(f'第{i}页',content)
                if not content.strip(): notes.append(f'第{i}页无可提取文字，可能需要OCR')
        notes.append('PDF表格列序可能变化；涉及参数表和图片须查看原页。')
    else: raise ValueError('不支持的格式')
    if not any(u['text'].strip() for u in units): notes.append('未提取到有效文本')
    return units,notes

def search(roots,terms,limit):
    needles=list(dict.fromkeys(norm(t.strip()) for t in terms if t.strip()))
    if not needles: raise ValueError('至少提供一个非空检索词')
    files,warnings,states=inventory(roots); found=[]; parsed=0; failed=0
    for f in files:
        if not f['supported']: continue
        try:
            units,notes=extract(resolve(roots,f['category'],f['file'])); parsed+=1
            hits=[(i,sum(t in norm(u['text']) for t in needles)) for i,u in enumerate(units)]
            hits=[(i,n) for i,n in hits if n]
            name_hits=sum(t in norm(f['file']) for t in needles)
            if hits or name_hits:
                excerpts=[{'unit':i+1,'locator':units[i]['locator'],'text':units[i]['text'][:1200],'truncated':len(units[i]['text'])>1200} for i,n in sorted(hits,key=lambda h:-h[1])[:5]]
                found.append({**f,'score':name_hits*3+sum(n for i,n in hits),'unit_count':len(units),'excerpts':excerpts,'extraction_notes':notes})
            for note in notes: warnings.append({'category':f['category'],'file':f['file'],'note':note})
        except Exception as e:
            failed+=1; warnings.append({'category':f['category'],'file':f['file'],'error':str(e)})
    found.sort(key=lambda f:(-f['score'],f['category'],f['file']))
    return {'folders':states,'scanned_files':len(files),'parsed_files':parsed,'failed_files':failed,'total_matches':len(found),'results':found[:limit],'results_truncated':len(found)>limit,'warnings':warnings,'notice':'关键词仅定位候选资料，需同义词补检并read完整条目，零命中不代表不存在。'}

def main():
    a=argparse.ArgumentParser(); a.add_argument('--config'); sub=a.add_subparsers(dest='cmd',required=True)
    cfg=sub.add_parser('configure'); cfg.add_argument('--root',required=True); cfg.add_argument('--replace',action='store_true')
    sub.add_parser('status'); sub.add_parser('inventory')
    q=sub.add_parser('search'); q.add_argument('--term',action='append',required=True); q.add_argument('--limit',type=int,default=15)
    r=sub.add_parser('read'); r.add_argument('--category',choices=CATS,required=True); r.add_argument('--file',required=True); r.add_argument('--start',type=int,default=1); r.add_argument('--count',type=int,default=80)
    args=a.parse_args(); cp=config_path(args.config)
    if args.cmd=='configure':
        root=Path(args.root).expanduser()
        if not root.is_absolute(): raise ValueError('请输入当前执行环境可访问的绝对路径')
        if cp.exists() and not args.replace: raise ValueError('配置已存在；修改路径请显式使用--replace')
        roots={k:root/k for k in CATS}; states=status(roots)
        if any(v['status']!='accessible' for v in states.values()): emit({'status':'not_configured','folders':states}); return 2
        cp.parent.mkdir(parents=True,exist_ok=True); cp.write_text(json.dumps({'schema_version':1,'folders':{k:str(v) for k,v in roots.items()}},ensure_ascii=False,indent=2),encoding='utf-8')
        emit({'status':'configured','config':str(cp),'folders':states}); return 0
    roots=load_config(cp)
    if args.cmd=='status': emit({'config':str(cp),'folders':status(roots)})
    elif args.cmd=='inventory':
        files,warnings,states=inventory(roots); emit({'folders':states,'files':files,'warnings':warnings})
    elif args.cmd=='search': emit(search(roots,args.term,max(1,min(args.limit,200))))
    elif args.cmd=='read':
        if args.start<1 or not 1<=args.count<=500: raise ValueError('start须>=1，count须为1至500')
        p=resolve(roots,args.category,args.file); units,notes=extract(p); sel=units[args.start-1:args.start-1+args.count]
        emit({'category':args.category,'file':args.file,'mtime_ns':p.stat().st_mtime_ns,'total_units':len(units),'units':[{'unit':args.start+i,**u} for i,u in enumerate(sel)],'has_more':args.start-1+len(sel)<len(units),'warnings':notes})
    return 0
if __name__=='__main__':
    if hasattr(sys.stdout,'reconfigure'): sys.stdout.reconfigure(encoding='utf-8')
    try: sys.exit(main())
    except Exception as e: emit({'status':'error','error':str(e),'notice':'读取或配置失败，不能据此判断企业知识不存在。'}); sys.exit(2)
